from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\micha\OneDrive\Documents\Sites\isu-proj-hub\manuals\ISU_ETS_New_Features_Navigation_Manual.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header_fill="1F4E79"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.size = Pt(10)


def add_feature_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(1.75)
    table.columns[2].width = Inches(2.35)
    table.columns[3].width = Inches(1.65)
    headers = ["Feature", "Where to Navigate", "What to Do", "Role Access"]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    for feature, where, action, role in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], feature)
        set_cell_text(cells[1], where)
        set_cell_text(cells[2], action)
        set_cell_text(cells[3], role)
    style_table(table)
    doc.add_paragraph()


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for style_name in ("Heading 1", "Heading 2", "Heading 3"):
    styles[style_name].font.name = "Aptos"
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("ISU ETS Project Hub New Features Navigation Manual")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("Guide for Admin and Faculty users")
run.font.size = Pt(12)
run.font.bold = True

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(10)
intro.add_run(
    "This manual explains where to find and how to use the newly added features in the "
    "ISU Cauayan Extension Training Services Project Management Hub. It is written for "
    "Admin users who review and monitor records, and for Faculty users who create and "
    "maintain assigned program, project, component, activity, participant, document, and "
    "proposal records."
)

add_heading(doc, "Quick Navigation Summary", 1)
summary_rows = [
    ("Total requests", "Dashboard", "Open the Total Requests card or Proposals filtered to submitted requests.", "Admin and Faculty"),
    ("Programs Projects Components Activities", "Project Management", "Use the tabs to move through the project hierarchy and create assigned records.", "Faculty can add within scope"),
    ("Participants", "Project Management > Activities", "Open an activity and click the Participants action to view or add participant records.", "Faculty add, Admin view"),
    ("Partner emails", "Partners", "Open a partner record and use Send Email to contact project partners.", "Admin"),
    ("Beneficiary counts", "Beneficiaries", "Review beneficiary groups and linked participant counts.", "Admin"),
    ("Document uploads", "Documents", "Upload scanned project files and required attachments.", "Admin and assigned Faculty"),
    ("Faculty timeline", "Faculty", "Open a faculty profile to review activity history and assignments.", "Admin"),
    ("Proposal remarks and history", "Proposals", "Open proposal details to read comments, budget remarks, files, and status history.", "Admin and Faculty"),
    ("Reports", "Accomplishments", "Filter by year or quarter and generate a report template.", "Admin and Faculty view"),
    ("Designations", "Project Management Assign action and Designation page", "Assign faculty, update dates, record accomplishments, and view designation files.", "Admin assign, Faculty view"),
    ("Certificates and attendance", "Certificates", "Import attendance and generate certificates from present participants.", "Admin generate, Faculty view"),
]
add_feature_table(doc, summary_rows)

add_heading(doc, "Faculty Workflow", 1)
add_bullets(doc, [
    "Start on Dashboard to check assigned programs, assigned projects, assigned activities, participants, total requests, and notifications.",
    "Open Project Management, then choose Programs, Projects, Components, or Activities from the tabs.",
    "Create or update records only within the assigned scope shown to the account.",
    "Open Activities and use the Participants action to monitor participants or add new participant records.",
    "Open Proposals to submit proposals, revise returned proposals, read remarks, upload required files, and check previous submission status.",
    "Open Documents to upload scanned project documents when the record is inside the assigned scope.",
    "Open Designation, MOA, Accomplishments, and Certificates to view, print, or download visible records.",
])

add_heading(doc, "Admin Workflow", 1)
add_bullets(doc, [
    "Start on Dashboard to monitor total requests, pending reviews, notifications, project progress, reports, and system activity.",
    "Open Project Management to view programs, projects, components, activities, files, and assignments.",
    "Use the Assign action on program, project, component, or activity rows to assign faculty designations.",
    "Open Proposals to review submissions, add general remarks, add budget remarks, approve, return, or reject proposals.",
    "Open Users to manage account information and review submission history for each user.",
    "Open Reports, Designation, Certificates, Documents, Partners, Beneficiaries, Faculty, Audit Logs, and Settings for monitoring and administrative maintenance.",
])

add_heading(doc, "New Feature Details", 1)

add_heading(doc, "Dashboard Cards and Notifications", 2)
add_bullets(doc, [
    "Navigate to Dashboard from the sidebar.",
    "Click Programs, Projects, Activities, Participants, or Total Requests cards to open the matching module.",
    "Use the bell icon in the top navigation to view proposal approvals, remarks, designation notices, and other system notifications.",
    "Email notifications are sent for proposal approvals and proposal remarks when mail settings are configured.",
])

add_heading(doc, "Activities and Participants", 2)
add_bullets(doc, [
    "Navigate to Project Management > Activities.",
    "Click New Activity to create an activity under an accessible component.",
    "Click the Participants icon on an activity row to list participants.",
    "Faculty users can add participants when the activity belongs to their accessible assigned scope.",
    "Admin users can view participant lists but are blocked from adding participants in the faculty participant workflow.",
])

add_heading(doc, "Partners and Beneficiaries", 2)
add_bullets(doc, [
    "Navigate to Partners to maintain partner agency records and send email messages to partners.",
    "Navigate to Beneficiaries to maintain beneficiary groups and view the number of linked beneficiaries or participants.",
    "Participant beneficiary group selections are used by activity participant records and beneficiary counts.",
])

add_heading(doc, "Documents and Scanned Files", 2)
add_bullets(doc, [
    "Navigate to Documents.",
    "Use Upload Document to attach scanned files or required project files.",
    "Choose the correct record type, such as program, project, component, or activity, so the file appears in the right workflow.",
    "Faculty can upload only within accessible assigned scope. Admin can manage broader document records.",
])

add_heading(doc, "Proposals Remarks and Submission History", 2)
add_bullets(doc, [
    "Navigate to Proposals.",
    "Faculty can add proposals, revise returned proposals, upload required files, and review status history.",
    "Admin can open submitted proposals, add remarks, add budget comments, approve, return, or reject.",
    "Open proposal details to see previous proposals, proposal status, comments, budget remarks, and attached files.",
])

add_heading(doc, "Reports and Accomplishments", 2)
add_bullets(doc, [
    "Navigate to Accomplishments.",
    "Use year filters to view previous-year reports, including 2021, 2022, 2023, and 2024.",
    "Use quarter filters to review Q1, Q2, Q3, and Q4 records.",
    "Click Generate Template to create a report template with proposal title, project leader or head, MOA, project report, activities, related project information, and partner agency requirements.",
])

add_heading(doc, "Designation Management", 2)
add_bullets(doc, [
    "Admin opens Project Management and clicks the Assign action on a program, project, component, or activity.",
    "Use the assignment panel to set assigned faculty, designation dates, and accomplishment notes.",
    "Faculty can open Designation to view visible designation files and can receive popup or dashboard notifications when assigned.",
    "Faculty access is view, print, or download only for designation files; upload is reserved for Admin.",
])

add_heading(doc, "Certificates and Attendance", 2)
add_bullets(doc, [
    "Navigate to Certificates.",
    "Admin can use Import Attendance to upload attendance CSV records.",
    "Admin can use Generate From Attendance to create certificates for present participants.",
    "Users can view and print generated certificates from the Certificates page.",
])

add_heading(doc, "User Records", 2)
add_bullets(doc, [
    "Admin navigates to Users.",
    "Use New User to create an account.",
    "Use the action buttons to view, edit, or delete users.",
    "Open user details to review basic user information and submission history.",
])

add_heading(doc, "Role Access Matrix", 1)
role_rows = [
    ("Project hierarchy records", "Create and update within assigned scope", "View all and assign faculty"),
    ("Participants", "Add within accessible activity scope", "View participant lists"),
    ("Reports", "View print download visible records", "Upload manage generate templates"),
    ("Designation files", "View print download", "Upload manage assign dates accomplishments"),
    ("MOA records", "View print visible records", "Manage and monitor"),
    ("Proposals", "Submit revise comment view history", "Review comment approve return reject"),
    ("Certificates", "View print visible records", "Import attendance generate certificates"),
    ("Users settings audit logs", "No access", "Full administrative access"),
]
table = doc.add_table(rows=1, cols=3)
for index, header in enumerate(["Area", "Faculty", "Admin"]):
    set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
for area, faculty, admin in role_rows:
    cells = table.add_row().cells
    set_cell_text(cells[0], area)
    set_cell_text(cells[1], faculty)
    set_cell_text(cells[2], admin)
style_table(table, header_fill="404040")

doc.add_section(WD_SECTION.NEW_PAGE)
add_heading(doc, "Recommended Testing Checklist", 1)
add_bullets(doc, [
    "Login as Admin and confirm Dashboard, Proposals, Reports, Designation, Certificates, Users, and Project Management pages load.",
    "Login as Faculty and confirm the sidebar hides Users and Settings.",
    "As Faculty, open Project Management and confirm Programs, Projects, Components, and Activities tabs are visible.",
    "As Faculty, open Activities and confirm participant actions are available on accessible activities.",
    "As Admin, confirm participant add attempts are blocked while participant lists remain viewable.",
    "Confirm faculty upload buttons are hidden on Designation and Reports where view or print access is required.",
    "Confirm proposal remarks and approval actions show loading feedback and refresh after completion.",
])

doc.save(OUTPUT)
print(OUTPUT)
